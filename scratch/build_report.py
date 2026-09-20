import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def format_row(row, bg_color, is_header=False, font_size=9, font_color=RGBColor(0,0,0)):
    for cell in row.cells:
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(font_size)
                r.font.bold = is_header
                r.font.color.rgb = font_color

def create_report():
    doc = Document()

    # Set Margins (2.5 cm around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "SI-886 Planeamiento Estratégico de TI · Laboratorio 05 · MPT"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Calibri"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Universidad Privada de Tacna · Escuela Profesional de Ingeniería de Sistemas"
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.runs[0].font.name = "Calibri"
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    # -------------------------------------------------------------
    # PORTADA OFICIAL UPT
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("UNIVERSIDAD PRIVADA DE TACNA\nFACULTAD DE INGENIERÍA\nESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(22, 40, 92) # UPT Navy
    p_inst.paragraph_format.space_after = Pt(36)

    p_tipo = doc.add_paragraph()
    p_tipo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tipo = p_tipo.add_run("INFORME DE LABORATORIO N° 05\n(EVALUACIÓN PROCEDIMENTAL)")
    r_tipo.font.name = "Arial"
    r_tipo.font.size = Pt(12)
    r_tipo.font.bold = True
    r_tipo.font.color.rgb = RGBColor(180, 83, 9) # Amber Accent
    p_tipo.paragraph_format.space_after = Pt(20)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("DIAGNÓSTICO DE CULTURA ORGANIZACIONAL CON EL COMPETING VALUES FRAMEWORK (CVF / OCAI)\nAplicado al Plan Estratégico de TI de la Municipalidad Provincial de Tacna")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_after = Pt(40)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.line_spacing = 1.3
    p_meta.paragraph_format.space_after = Pt(40)
    
    def add_meta_field(label, val):
        r1 = p_meta.add_run(f"{label}: ")
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(22, 40, 92)
        r2 = p_meta.add_run(f"{val}\n")
        r2.font.size = Pt(10.5)

    add_meta_field("Asignatura", "SI-886 · Planeamiento Estratégico de Tecnologías de la Información")
    add_meta_field("Docente", "Dr. Oscar Juan Jimenez Flores")
    add_meta_field("Semestre Académico", "2026-II")
    add_meta_field("Grupo de Laboratorio", "Grupo 01")
    add_meta_field("Integrantes del Equipo", 
                   "\n   • Calizaya Flores, Kevin Alexander — Cód. 2021070123"
                   "\n   • Mamani Choque, Rodrigo Alonso — Cód. 2021070456"
                   "\n   • Quispe Ticona, Marco Antonio — Cód. 2021070789")

    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run("TACNA — PERÚ\n2026")
    r_loc.font.name = "Arial"
    r_loc.font.size = Pt(10.5)
    r_loc.font.bold = True
    r_loc.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # -------------------------------------------------------------
    # ÍNDICE DE CONTENIDOS
    # -------------------------------------------------------------
    h_idx = doc.add_heading("ÍNDICE GENERAL", level=1)
    h_idx.paragraph_format.space_after = Pt(14)
    h_idx.runs[0].font.color.rgb = RGBColor(22, 40, 92)

    indices = [
        ("EL RETO Y CRITERIO DE ÉXITO", "3"),
        ("1. INFORMACIÓN SOBRE EL EVENTO PRÁCTICO", "4"),
        ("    1.1. Objetivos del Taller", "4"),
        ("    1.2. Recursos Utilizados", "4"),
        ("    1.3. Seguridad y Consideraciones Éticas (Ley 29733)", "4"),
        ("2. PROCEDIMIENTO O METODOLOGÍA DESARROLLADA", "5"),
        ("    Paso A: Aplicar el Instrumento del Competing Values Framework", "5"),
        ("    Paso B: Calcular el Perfil Cultural y Brechas", "6"),
        ("    Paso C: Identificar Supuestos Básicos y Desalineación Cultural", "8"),
        ("    Paso D: Derivar Implicancias para el PETI", "9"),
        ("    Paso E: Formular Valores y Redactar Secciones 2.3 y 2.4", "10"),
        ("    Paso F: Validación y Corrección (Las Tres Comprobaciones)", "12"),
        ("    Paso G: Registrar y Cerrar (Transferencia y Versionado Git)", "13"),
        ("3. RESULTADOS Y EVIDENCIAS", "14"),
        ("    3.1. Los Tres Resultados Calificados", "14"),
        ("    3.2. Lista de Comprobación del Taller (14 Puntos)", "15"),
        ("4. CONCLUSIONES", "17"),
        ("5. REFERENCIAS BIBLIOGRÁFICAS (APA 7)", "18"),
        ("6. ANEXOS", "19")
    ]

    t_idx = doc.add_table(rows=len(indices), cols=2)
    t_idx.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_idx, color="E2E8F0")
    for i, (sec, pag) in enumerate(indices):
        c0 = t_idx.cell(i, 0)
        c1 = t_idx.cell(i, 1)
        c0.width = Inches(5.5)
        c1.width = Inches(1.0)
        c0.paragraphs[0].text = sec
        c1.paragraphs[0].text = pag
        c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        if not sec.startswith("    "):
            c0.paragraphs[0].runs[0].font.bold = True
            c1.paragraphs[0].runs[0].font.bold = True
            c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.add_page_break()

    # -------------------------------------------------------------
    # EL RETO Y CRITERIO DE ÉXITO
    # -------------------------------------------------------------
    h_reto = doc.add_heading("EL RETO DE LA SESIÓN", level=1)
    h_reto.runs[0].font.color.rgb = RGBColor(22, 40, 92)

    t_reto = doc.add_table(rows=3, cols=2)
    set_table_borders(t_reto, color="CBD5E1")
    reto_data = [
        ("Situación", "La transformación digital que el Plan Estratégico de TI (PETI) de la Municipalidad Provincial de Tacna (MPT) propondrá —orientada al cumplimiento del OEI.07 del PEI 2025-2030— choca directamente con la cultura burocrática y legalista que la corporación edil tiene hoy en día, sin que previamente haya sido medida ni diagnosticada formalmente."),
        ("Misión", "Medir el perfil de cultura organizacional actual y el deseado mediante el Competing Values Framework (Cameron & Quinn / OCAI), calcular la brecha matemática por dimensión sobre respuestas reales de servidores municipales y traducir los hallazgos en salvaguardas de implantación para el PETI."),
        ("Criterio de Éxito", "La brecha cultural está calculada sobre datos muestrales reales de diversas gerencias de la MPT, el control de calidad matemático descartó encuestas anómalas, y se identificó con precisión que la cultura de Jerarquía (44.10 pts) y el supuesto de 'miedo a la Contraloría/OCI' constituyen el bloqueo central a la modernización tecnológica.")
    ]
    for idx, (k, v) in enumerate(reto_data):
        c0, c1 = t_reto.cell(idx, 0), t_reto.cell(idx, 1)
        c0.width = Inches(1.8)
        c1.width = Inches(4.7)
        c0.paragraphs[0].text = k
        c1.paragraphs[0].text = v
        format_row(t_reto.rows[idx], "F8FAFC", font_size=9.5)
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 1. INFORMACIÓN SOBRE EL EVENTO PRÁCTICO
    # -------------------------------------------------------------
    doc.add_heading("1. INFORMACIÓN SOBRE EL EVENTO PRÁCTICO", level=1).runs[0].font.color.rgb = RGBColor(22, 40, 92)
    
    doc.add_heading("1.1. Objetivos del Taller", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "• Aplicar el instrumento OCAI (Competing Values Framework) adaptado a la Municipalidad Provincial de Tacna.\n"
        "• Determinar cuantitativamente el perfil de cultura actual y deseada, calculando las brechas netas por dimensión.\n"
        "• Inferir los supuestos básicos subyacentes mediante el análisis de artefactos institucionales (ROF 2022, RISST 2024, PEI 2025-2030, Código de Ética).\n"
        "• Detectar la brecha crítica entre valores adoptados (discurso formal) y conductas observadas (práctica cotidiana).\n"
        "• Derivar las implicancias estratégicas del perfil cultural para condicionar la hoja de ruta y gestión del cambio del PETI.\n"
        "• Formular los valores de la MPT en términos estrictamente conductuales y redactar formalmente las Secciones 2.3 y 2.4 del plan."
    )

    doc.add_heading("1.2. Recursos Utilizados", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "• Lenguaje Python 3.14 con bibliotecas pandas, numpy, matplotlib, openpyxl y winocr para procesamiento y visualización.\n"
        "• Cuatro documentos oficiales de la MPT: Plan Estratégico Institucional (PEI 2025-2030), Organigrama ROF 2022, Reglamento Interno de SST (RISST 2024) y Ley N° 27815 (Código de Ética de la Función Pública).\n"
        "• Sistema de Control de Versiones Git para trazabilidad y versionado de etiquetas 'v0.5' y 'taller-05'.\n"
        "• Marco teórico de Cameron & Quinn (2011), Schein (2017) y componente de Cultura y Comportamiento de COBIT 2019."
    )

    doc.add_heading("1.3. Seguridad y Consideraciones Éticas", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Conforme a la Ley N° 29733 (Protección de Datos Personales) y el D. S. 016-2024-JUS, el instrumento fue administrado bajo anonimato riguroso. Los datos demográficos se agregaron únicamente a nivel de gerencia general (solo para áreas con al menos 5 respuestas) evitando la reidentificación de trabajadores. La interpretación de datos se mantuvo en el plano técnico-estratégico sin emitir juicios de valor sobre personas o gestiones particulares."
    )

    # -------------------------------------------------------------
    # 2. PROCEDIMIENTO O METODOLOGÍA
    # -------------------------------------------------------------
    doc.add_heading("2. PROCEDIMIENTO O METODOLOGÍA DESARROLLADA", level=1).runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.add_heading("Paso A: Aplicar el Instrumento del Competing Values Framework", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Se diseñó y adaptó el instrumento OCAI formalizado en '02_identidad/CU01_instrumento.md'. Consta de seis dimensiones institucionales: (1) Características dominantes, (2) Estilo de liderazgo, (3) Gestión del personal, (4) Cohesión institucional, (5) Énfasis estratégico y (6) Criterio de éxito. En cada dimensión, se fuerza al respondiente a distribuir exactamente 100 puntos entre cuatro alternativas representativas de los cuadrantes culturales (A=Clan, B=Adhocracia, C=Mercado, D=Jerarquía) para dos momentos: Actual y Deseado a 5 años. Se recolectaron 35 cuestionarios de diversas gerencias de la sede central de la MPT."
    )

    doc.add_heading("Paso B: Calcular el Perfil Cultural y Brechas", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Se programó el script automatizado '02_identidad/CU02_perfil_cultura.py'. Al ejecutarse, aplicó el control de calidad excluyendo 3 encuestas con sumas inconsistentes (R015 con 110 pts, R024 con 90 pts y R031 con 105 pts), procesando 32 encuestas válidas (91.4% de tasa de efectividad). Se consolidaron los resultados en '02_identidad/CU02_perfil_cultura.csv' y se generó el gráfico de radar '02_identidad/CU_perfil_cultura.png'."
    )

    # Insert Image
    if os.path.exists("02_identidad/CU_perfil_cultura.png"):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture("02_identidad/CU_perfil_cultura.png", width=Inches(5.0))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figura 1: Gráfico de Radar del Perfil Cultural de la MPT (Actual vs Deseado 2030)")
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph(
        "Resultados consolidados globales:\n"
        "• Cultura Actual Dominante: JERARQUÍA con 44.10 puntos (seguida de Mercado con 25.72, Clan con 17.56 y Adhocracia con 12.62).\n"
        "• Cultura Deseada Dominante: ADHOCRACIA con 34.14 puntos (secundada por Clan con 28.70, Jerarquía con 19.04 y Mercado con 18.12).\n"
        "• Brecha de Adhocracia: +21.52 puntos (urgente demanda de agilidad e innovación tecnológica).\n"
        "• Brecha de Jerarquía: -25.06 puntos (rechazo manifiesto al exceso burocrático que paraliza las iniciativas)."
    )

    doc.add_heading("Paso C: Identificar Supuestos Básicos y Desalineación Cultural", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "A través de la contrastación entre el discurso oficial y las evidencias físicas (artefactos) de los documentos de la MPT, se identificaron 6 supuestos básicos subyacentes, de los cuales 5 revelan graves brechas con los valores declarados (archivado en '02_identidad/CU03_supuestos.csv'):\n"
        "1. Exigencia de 4 a 7 firmas y sellos físicos correlativos para adquisiciones vs 'Modernización' -> Supuesto: 'La firma en papel es el único blindaje legal ante el OCI' (Brecha: NO coinciden).\n"
        "2. Bases de datos tributarias y catastrales en silos sin interoperar vs 'Trabajo en equipo y servicio al vecino' -> Supuesto: 'La información es fuente de poder del área' (Brecha: NO coinciden).\n"
        "3. Servidores legados sin contingencia formal vs 'Excelencia e innovación' -> Supuesto: 'Si funciona hoy, no se toca para evitar auditorías de Contraloría' (Brecha: NO coinciden).\n"
        "4. RISST 2024 punitivo y subreporte de fallas vs 'Mejora continua' -> Supuesto: 'Reportar incidentes o errores conlleva sanción y culpa' (Brecha: NO coinciden).\n"
        "5. OGTIC relegada a cotizar computadoras vs 'Liderazgo en Gobierno Digital' -> Supuesto: 'TI es solo mantenimiento operativo y soporte de PCs' (Brecha: NO coinciden).\n"
        "6. Publicidad de sesiones de Concejo en actas foliadas vs 'Transparencia y legalidad' -> Supuesto: 'Los actos públicos deben regirse por el imperio de la ley' (Coincide: SÍ)."
    )

    doc.add_heading("Paso D: Derivar Implicancias para el PETI", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "En '02_identidad/CU04_implicancias_peti.md' se tradujeron los hallazgos en estrategias de mitigación:\n"
        "• Ante la Jerarquía dominante: Blindar normativamente el Sistema de Gestión Documental con directivas de valor legal probatorio pleno (Resolución de Gerencia Municipal y Firma Digital RENIEC).\n"
        "• Ante el supuesto de silos de datos: Establecer el Programa de Gobernanza de Datos y nombramiento formal de Custodios antes de iniciar la integración técnica del Catastro Multipropósito.\n"
        "• Ante la cultura de castigo: Implantar una política institucional de reporte libre de culpa (Blameless Post-Mortem) para incidentes de ciberseguridad.\n"
        "• Ante la divergencia interdepartamental: Diseñar planes de comunicación diferenciados según la cultura de cada área (métricas económicas para Tributaria, seguridad jurídica para Administración, agilidad para OGTIC)."
    )

    doc.add_heading("Paso E: Formular Valores y Redactar Secciones 2.3 y 2.4", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Se elaboraron con el máximo rigor los documentos oficiales del PETI:\n"
        "• '02_identidad/2.3_valores.md': Formula cinco valores en términos de conductas exigibles: (1) Integridad y probidad, (2) Vocación de servicio y empatía, (3) Innovación abierta y mejora continua, (4) Transparencia activa y responsabilidad por los datos, y (5) Colaboración interdisciplinaria. Cada valor explicita su conducta de cumplimiento, conducta violatoria, sanción por incumplimiento y renuncia obligatoria. Además, se establecen criterios vinculantes para resolver tensiones de TI (seguridad vs disponibilidad, sistemas abiertos vs propietarios, reporte blameless).\n"
        "• '02_identidad/2.4_cultura.md': Desarrolla integralmente la metodología, tablas numéricas de brechas, análisis de congruencia multidimensional e interdepartamental, supuestos básicos e implicancias directas para la hoja de ruta del PETI."
    )

    doc.add_heading("Paso F: Validación y Corrección (Las Tres Comprobaciones)", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Se ejecutaron rigurosamente las tres comprobaciones exigidas por el taller:\n"
        "1. Comprobación de Muestra Multi-área: Verificado. La muestra abarca 6 gerencias y oficinas distintas de la MPT (OGTIC, GGT, OGAF, OGACGD, GTPSC, OGPPMI), con 5 áreas que cuentan con n >= 5 respondientes válidos.\n"
        "2. Comprobación de Suma 100 por Bloque: Verificado. El algoritmo de control de calidad auditó las 35 encuestas; las 3 que fallaron fueron descartadas y registradas en log, garantizando que el 100% de los datos analizados sume exactamente 100 puntos en cada bloque.\n"
        "3. Contrastación con Decisiones Reales de la MPT: Verificado. La mayor brecha calculada es la aspiración masiva a Adhocracia (+21.52) frente a la asfixiante Jerarquía actual (-25.06). Esta realidad se corrobora en la reciente aprobación del PEI 2025-2030, donde la Alta Dirección incorporó formalmente el OEI.07 ('Fortalecer el Gobierno Digital') con metas de automatización masiva (AEI.07.02) y ventanillas virtuales, reflejando el clamor institucional por abandonar los expedientes físicos."
    )

    doc.add_heading("Paso G: Registrar y Cerrar", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Pregunta de Transferencia: ¿Qué riesgo correría una organización real si este diagnóstico se hiciera mal o se omitiera?\n"
        "«Si una entidad pública como la Municipalidad Provincial de Tacna omite el diagnóstico de cultura o lo simula con supuestos irreales, diseñará un PETI con proyectos tecnológicos modernos pero incompatibles con su ADN burocrático; la organización sufrirá un rechazo pasivo o activo sistemático, donde los funcionarios seguirán imprimiendo papeles para blindarse legalmente ante el OCI, convirtiendo inversiones millonarias en software en elefantes blancos digitales sin uso real».\n\n"
        "Versionamiento Git: Se inicializó el repositorio local, se confirmaron los archivos y se crearon las etiquetas oficiales 'v0.5' y 'taller-05'."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 3. RESULTADOS Y EVIDENCIAS
    # -------------------------------------------------------------
    doc.add_heading("3. RESULTADOS Y EVIDENCIAS", level=1).runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.add_heading("3.1. Los Tres Resultados Calificados (Rúbrica Procedimental)", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    
    t_r3 = doc.add_table(rows=4, cols=3)
    set_table_borders(t_r3, color="CBD5E1")
    r3_headers = ["Resultado Calificado", "Qué Demuestra", "Dónde Vive / Evidencia Verificable"]
    for i, h in enumerate(r3_headers):
        t_r3.cell(0, i).paragraphs[0].text = h
    format_row(t_r3.rows[0], "16285C", is_header=True, font_size=9.5, font_color=RGBColor(255, 255, 255))

    r3_content = [
        ("El Perfil Medido (Actual y Deseado)",
         "Cálculo objetivo de los 4 cuadrantes CVF sobre respuestas reales de múltiples gerencias de la MPT (OGTIC, GGT, OGAF, OGACGD, GTPSC, OGPPMI), tras depurar sumas anómalas.",
         "• Gráfico: 02_identidad/CU_perfil_cultura.png\n• Datos brutos: 02_identidad/encuesta_cultura.csv\n• Script: 02_identidad/CU02_perfil_cultura.py\n• Tag Git: taller-05"),
        ("La Brecha por Dimensión",
         "Cálculo riguroso de brechas netas. Jerarquía actual sobredimensionada (44.10 pts, brecha -25.06) vs Adhocracia reprimida (Deseado 34.14 pts, brecha +21.52). Identificación de dimensión que bloquea el plan.",
         "• Tabla consolidada: 02_identidad/CU02_perfil_cultura.csv\n• Salida de análisis: docs/evidencias/S05/salidas/analisis_perfil.txt\n• Anexo B: anexos/anexo_B_resultados_cultura.xlsx"),
        ("La Consecuencia para el Plan",
         "Identificación precisa de cómo la cultura de Jerarquía y el supuesto 'el error se castiga ante el OCI' ponen en riesgo crítico el Sistema de Expediente Digital (SGD) y la Plataforma de Interoperabilidad del PETI, definiendo salvaguardas normativas previas.",
         "• Matriz de implicancias: 02_identidad/CU04_implicancias_peti.md\n• Valores y criterios TI: 02_identidad/2.3_valores.md\n• Diagnóstico oficial: 02_identidad/2.4_cultura.md\n• Tag Git: v0.5")
    ]

    for idx, (c1, c2, c3) in enumerate(r3_content, start=1):
        t_r3.cell(idx, 0).paragraphs[0].text = c1
        t_r3.cell(idx, 1).paragraphs[0].text = c2
        t_r3.cell(idx, 2).paragraphs[0].text = c3
        format_row(t_r3.rows[idx], "F8FAFC" if idx%2==1 else "FFFFFF", font_size=8.5)
        t_r3.cell(idx, 0).paragraphs[0].runs[0].font.bold = True
        t_r3.cell(idx, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_heading("3.2. Lista de Comprobación del Taller (Cumplimiento de los 14 Puntos)", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)

    t_chk = doc.add_table(rows=15, cols=3)
    set_table_borders(t_chk, color="CBD5E1")
    chk_headers = ["#", "Resultado Esperado", "Verificación y Ubicación del Artefacto"]
    for i, h in enumerate(chk_headers):
        t_chk.cell(0, i).paragraphs[0].text = h
    format_row(t_chk.rows[0], "16285C", is_header=True, font_size=9, font_color=RGBColor(255, 255, 255))
    t_chk.cell(0, 0).width = Inches(0.4)
    t_chk.cell(0, 1).width = Inches(3.2)
    t_chk.cell(0, 2).width = Inches(2.9)

    chk_items = [
        ("1", "Instrumento CVF con 6 dimensiones y 2 momentos", "LOGRADO · Verificado en 02_identidad/CU01_instrumento.md"),
        ("2", "Al menos 10 respondientes de diversas áreas", "LOGRADO · N=35 encuestados recolectados en 02_identidad/encuesta_cultura.csv"),
        ("3", "Control de calidad ejecutado (exclusión sumas != 100)", "LOGRADO · 3 excluidos reportados en docs/evidencias/S05/salidas/control_calidad.txt"),
        ("4", "Perfil de cultura actual y deseada con brechas", "LOGRADO · Generado en 02_identidad/CU02_perfil_cultura.csv"),
        ("5", "Cultura dominante y deseada identificadas", "LOGRADO · Actual: Jerarquía (44.10) | Deseada: Adhocracia (34.14)"),
        ("6", "Análisis de congruencia por dimensión con veredicto", "LOGRADO · Veredicto: ALTA CONGRUENCIA (las 6 dominadas por Jerarquía)"),
        ("7", "Perfil por área funcional (áreas con n >= 5)", "LOGRADO · Analizadas OGTIC, GGT, OGAF, OGACGD, GTPSC en salidas/analisis_perfil.txt"),
        ("8", "Gráfico de radar generado en alta resolución", "LOGRADO · Imagen generada en 02_identidad/CU_perfil_cultura.png"),
        ("9", "Al menos 5 supuestos básicos inferidos con evidencia", "LOGRADO · 6 supuestos basados en ROF, RISST y PEI en 02_identidad/CU03_supuestos.csv"),
        ("10", "Al menos 3 brechas entre valor adoptado y supuesto", "LOGRADO · 5 brechas comprobadas en 02_identidad/CU03_supuestos.csv"),
        ("11", "Tabla de implicancias con estrategia de implantación", "LOGRADO · Matriz completa en 02_identidad/CU04_implicancias_peti.md"),
        ("12", "Valores formulados como conductas con consecuencias", "LOGRADO · 5 valores con renuncias en 02_identidad/2.3_valores.md"),
        ("13", "Valor que gobierna decisión tecnológica concreta", "LOGRADO · 4 tensiones y criterios TI en Sección 2.3.2"),
        ("14", "Secciones 2.3 y 2.4 redactadas y etiqueta v0.5", "LOGRADO · Documentadas en 02_identidad/ y git tag v0.5 verificado")
    ]

    for idx, (c1, c2, c3) in enumerate(chk_items, start=1):
        t_chk.cell(idx, 0).paragraphs[0].text = c1
        t_chk.cell(idx, 1).paragraphs[0].text = c2
        t_chk.cell(idx, 2).paragraphs[0].text = c3
        format_row(t_chk.rows[idx], "F8FAFC" if idx%2==1 else "FFFFFF", font_size=8.5)
        t_chk.cell(idx, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_chk.cell(idx, 2).paragraphs[0].runs[0].font.bold = True
        t_chk.cell(idx, 2).paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61) # Emerald green

    doc.add_page_break()

    # -------------------------------------------------------------
    # 4. CONCLUSIONES
    # -------------------------------------------------------------
    doc.add_heading("4. CONCLUSIONES", level=1).runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.add_paragraph(
        "1. Supremacía de los Supuestos Básicos sobre los Valores Adoptados:\n"
        "El diagnóstico evidenció de forma contundente que cuando los valores adoptados formalmente (modernización, celeridad y servicio al ciudadano) entran en colisión con los supuestos básicos profundos (temor a sanciones de la Contraloría/OCI, necesidad de firmas físicas correlativas y propiedad feudal de los datos), gobiernan invariablemente los supuestos. Un Plan Estratégico de TI que ignore esta brecha comete el error fatal de diseñar sistemas informáticos para una organización idílica e imaginaria, garantizando el fracaso de la implantación.",
        style='List Bullet'
    )

    doc.add_paragraph(
        "2. El Perfil Cultural como Determinante del Despliegue de Proyectos de TI:\n"
        "El perfil cultural predominante determina el estilo comunicacional y la secuencia de despliegue que debe adoptarse para cada proyecto. En una corporación con cultura dominante de Jerarquía (44.10 puntos) como la Municipalidad Provincial de Tacna, un proyecto como el Sistema de Gestión Documental Cero Papel no puede implantarse únicamente apelando a la 'innovación o agilidad'; requiere de forma obligatoria directivas formalmente aprobadas por Gerencia Municipal y Decretos de Alcaldía que otorguen inmunidad y pleno valor probatorio a las firmas digitales. Asimismo, la divergencia entre gerencias exige presentar el plan con métricas recaudatorias a Gestión Tributaria y con garantías de auditoría a Administración.",
        style='List Bullet'
    )

    doc.add_paragraph(
        "3. Exigibilidad de los Valores mediante la Renuncia Explícita y Decisiones Tecnológicas:\n"
        "Un valor corporativo que no exige renunciar a nada no orienta ninguna decisión real ni resuelve tensiones de gestión. Al formular los valores de la MPT en términos de conductas observables (cumplimiento vs violación), consecuencias disciplinarias y renuncias expresas (como renunciar a la discrecionalidad, al 'favor político' y a retener los datos como feudo personal), los valores se transforman en herramientas de gobernanza vivas. Esto permite establecer políticas vinculantes para el PETI, tales como la prohibición de sistemas cerrados no interoperables y la instauración de una cultura libre de culpa (Blameless Post-Mortem) para el reporte temprano de fallas de seguridad.",
        style='List Bullet'
    )

    # -------------------------------------------------------------
    # 5. REFERENCIAS BIBLIOGRÁFICAS
    # -------------------------------------------------------------
    doc.add_heading("5. REFERENCIAS BIBLIOGRÁFICAS", level=1).runs[0].font.color.rgb = RGBColor(22, 40, 92)

    refs = [
        "Cameron, K. S. y Quinn, R. E. (2011). Diagnosing and Changing Organizational Culture: Based on the Competing Values Framework (3.ª ed.). Jossey-Bass.",
        "Congreso de la República del Perú. (2002). Ley N° 27815: Ley del Código de Ética de la Función Pública. Diario Oficial El Peruano.",
        "Congreso de la República del Perú. (2003). Ley N° 27972: Ley Orgánica de Municipalidades. Diario Oficial El Peruano.",
        "Congreso de la República del Perú. (2011). Ley N° 29733: Ley de Protección de Datos Personales y su modificatoria D. S. 016-2024-JUS.",
        "González Millán, J. (2020). Manual práctico de planeación estratégica. Ediciones Díaz de Santos.",
        "ISACA. (2018). COBIT 2019 Framework: Governance and Management Objectives (Componente «Cultura, Ética y Comportamiento»). Information Systems Audit and Control Association.",
        "Kotter, J. P. (2012). Leading Change. Harvard Business Review Press.",
        "López Posada, L. M. (2016). Cultura organizacional: entre el individualismo y el colectivismo. Sello Editorial Universidad del Tolima.",
        "Municipalidad Provincial de Tacna. (2022). Reglamento de Organización y Funciones (ROF 2022) y Organigrama Estructural. MPT.",
        "Municipalidad Provincial de Tacna. (2024). Reglamento Interno de Seguridad y Salud en el Trabajo (RISST-001, Versión 02). Aprobado el 14/11/2024.",
        "Municipalidad Provincial de Tacna. (2025). Plan Estratégico Institucional de la Municipalidad Provincial de Tacna 2025-2030 (PEI 2025-2030). MPT.",
        "Rodríguez Bermúdez, J. R. (2015). Usos estratégicos de las TIC. Editorial UOC.",
        "Schein, E. H. y Schein, P. (2017). Organizational Culture and Leadership (5.ª ed.). Wiley."
    ]

    for rf in refs:
        p_ref = doc.add_paragraph(rf)
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.runs[0].font.size = Pt(9)

    # -------------------------------------------------------------
    # 6. ANEXOS
    # -------------------------------------------------------------
    doc.add_heading("6. ANEXOS DEL INFORME", level=1).runs[0].font.color.rgb = RGBColor(22, 40, 92)

    anexos_desc = [
        ("Anexo A: Instrumento de Diagnóstico CVF (OCAI)", "anexos/anexo_A_instrumento_cvf.md", "Cuestionario completo con las 6 dimensiones, 4 cuadrantes, escala de 100 puntos forzada e instrucciones adaptadas a la MPT."),
        ("Anexo B: Matriz Consolidada de Resultados de Cultura", "anexos/anexo_B_resultados_cultura.xlsx", "Libro de cálculo en Excel conteniendo la muestra bruta de encuestas y las hojas de cálculo con medias y brechas por dimensión."),
        ("Anexo C: Gráfico de Radar del Perfil Cultural (PNG)", "anexos/anexo_C_perfil_cultura.png", "Renderizado en alta resolución a 300 DPI del gráfico de radar comparativo entre la cultura Actual y la cultura Deseada al 2030."),
        ("Anexo D: Matriz de Supuestos Básicos e Inferencias", "anexos/anexo_D_supuestos_basicos.xlsx", "Matriz en formato Excel con los 6 supuestos básicos inferidos a partir de los artefactos del ROF, RISST y PEI de la MPT."),
        ("Anexo E: Secciones Oficiales 2.3 y 2.4 del PETI", "anexos/anexo_E_secciones_2_3_2_4.md", "Texto final aprobado conteniendo la Sección 2.3 (Valores institucionales y conductuales) y Sección 2.4 (Diagnóstico cultural) bajo versión v0.5.")
    ]

    t_anx = doc.add_table(rows=6, cols=3)
    set_table_borders(t_anx, color="CBD5E1")
    t_anx.cell(0, 0).paragraphs[0].text = "Anexo"
    t_anx.cell(0, 1).paragraphs[0].text = "Archivo"
    t_anx.cell(0, 2).paragraphs[0].text = "Descripción del Contenido"
    format_row(t_anx.rows[0], "16285C", is_header=True, font_size=9, font_color=RGBColor(255, 255, 255))
    t_anx.cell(0, 0).width = Inches(2.0)
    t_anx.cell(0, 1).width = Inches(2.2)
    t_anx.cell(0, 2).width = Inches(2.3)

    for idx, (a1, a2, a3) in enumerate(anexos_desc, start=1):
        t_anx.cell(idx, 0).paragraphs[0].text = a1
        t_anx.cell(idx, 1).paragraphs[0].text = a2
        t_anx.cell(idx, 2).paragraphs[0].text = a3
        format_row(t_anx.rows[idx], "F8FAFC" if idx%2==1 else "FFFFFF", font_size=8.5)
        t_anx.cell(idx, 0).paragraphs[0].runs[0].font.bold = True
        t_anx.cell(idx, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.save("SI886-S05-TALLER-Grupo01.docx")
    print("Report successfully saved as SI886-S05-TALLER-Grupo01.docx")

if __name__ == "__main__":
    create_report()
