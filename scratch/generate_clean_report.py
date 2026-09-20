"""
generate_clean_report.py
Genera el informe oficial en base a SI886-PLANTILLA-TALLER (1).docx:
- Elimina Paso G de la Sección 2 (Procedimiento).
- Elimina los cuadros pesados de la Sección 2 para dejar una redacción limpia y ágil.
- Mantiene la Figura 2.1 (Radar de cultura en alta resolución).
- Rellena completamente el cuadro oficial de Resultados (Table 1) con los 14 resultados de la guía.
- Actualiza el TOC con Word COM y exporta a PDF.
"""

import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def format_row(row, bg_color, is_header=False, font_size=8.5, font_color=RGBColor(0,0,0)):
    for cell in row.cells:
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(font_size)
                r.font.bold = is_header
                r.font.color.rgb = font_color

def insert_p_after(ref_p, text="", style="Normal", space_after=4, bold_prefix="", bold_color=RGBColor(22,40,92)):
    new_p_elm = OxmlElement("w:p")
    ref_p._p.addnext(new_p_elm)
    p = docx.text.paragraph.Paragraph(new_p_elm, ref_p._parent)
    try:
        p.style = style
    except:
        p.style = "Normal"
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = bold_color
    if text:
        r_t = p.add_run(text)
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(9.5)
    return p

def run():
    template_path = "SI886-PLANTILLA-TALLER (1).docx"
    output_docx = "SI886-S05-TALLER-Grupo01.docx"
    output_pdf = "SI886-S05-TALLER-Grupo01.pdf"

    shutil.copyfile(template_path, output_docx)
    doc = Document(output_docx)

    GH_BASE = "https://github.com/GianfrancoArocutipa/Taller05_ArocutipaGIna"
    GH_TAG = f"{GH_BASE}/tree/taller-05"

    # 1. ACTUALIZAR CARÁTULA
    doc.paragraphs[11].text = "“TALLER 05: DIAGNÓSTICO DE CULTURA CON EL COMPETING VALUES FRAMEWORK (CVF / OCAI)”"
    doc.paragraphs[11].runs[0].font.name = "Arial"
    doc.paragraphs[11].runs[0].font.size = Pt(12)
    doc.paragraphs[11].runs[0].font.bold = True
    doc.paragraphs[11].runs[0].font.color.rgb = RGBColor(22, 40, 92)

    doc.paragraphs[14].text = "“SI-886 · PLANEAMIENTO ESTRATÉGICO DE TECNOLOGÍAS DE LA INFORMACIÓN”"
    doc.paragraphs[14].runs[0].font.name = "Arial"
    doc.paragraphs[14].runs[0].font.size = Pt(10.5)
    doc.paragraphs[14].runs[0].font.bold = True

    doc.paragraphs[17].text = "AROCUTIPA AROCUTIPA, GIANFRANCO · CÓDIGO: 2023076790"
    doc.paragraphs[17].runs[0].font.name = "Calibri"
    doc.paragraphs[17].runs[0].font.size = Pt(10.5)
    doc.paragraphs[17].runs[0].font.bold = True

    doc.paragraphs[27].text = "TACNA — PERÚ\n2026"
    doc.paragraphs[27].runs[0].font.name = "Arial"
    doc.paragraphs[27].runs[0].font.size = Pt(10)
    doc.paragraphs[27].runs[0].font.bold = True

    # Ajustar espacios vacíos en carátula para evitar salto de página
    for idx in [25, 24, 23, 8, 7, 6]:
        p = doc.paragraphs[idx]
        if not p.text.strip():
            p_elem = p._p
            if p_elem.getparent() is not None:
                p_elem.getparent().remove(p_elem)

    # 2. ELIMINAR LA PÁGINA "CÓMO SE USA ESTA PLANTILLA"
    paragraphs_to_remove = []
    found_heading = False
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and "Cómo se usa esta plantilla" in p.text:
            found_heading = True
        if found_heading:
            if p.style.name == "Heading 1" and "1. Información" in p.text:
                break
            paragraphs_to_remove.append(p)

    for p in paragraphs_to_remove:
        p_elem = p._p
        if p_elem.getparent() is not None:
            p_elem.getparent().remove(p_elem)

    # 3. IDENTIFICAR ENCABEZADOS REALES (NO DEL TOC)
    p_dict = {}
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "\t" in txt:
            continue
        if txt.startswith("1.1"):
            p_dict["1.1"] = p
        elif txt.startswith("1.2"):
            p_dict["1.2"] = p
        elif txt.startswith("1.3"):
            p_dict["1.3"] = p
        elif txt.startswith("1.4"):
            p_dict["1.4"] = p
        elif txt.startswith("1.5"):
            p_dict["1.5"] = p
        elif txt.startswith("1.6"):
            p_dict["1.6"] = p
        elif txt.startswith("2."):
            p_dict["2"] = p
        elif txt == "Paso A":
            p_dict["Paso A"] = p
        elif txt == "Paso B":
            p_dict["Paso B"] = p
        elif txt == "Paso C":
            p_dict["Paso C"] = p
        elif txt.startswith("3."):
            p_dict["3"] = p
        elif txt.startswith("4."):
            p_dict["4"] = p
        elif txt.startswith("5."):
            p_dict["5"] = p
        elif txt.startswith("6."):
            p_dict["6"] = p
        elif txt.startswith("7."):
            p_dict["7"] = p

    # 1.1 Título del evento práctico
    cur = p_dict["1.1"]
    cur = insert_p_after(cur, "Taller de Laboratorio 05 · Diagnóstico de cultura con el Competing Values Framework (Caso: Municipalidad Provincial de Tacna - MPT)", style="Normal", space_after=6)

    # 1.2 Objetivos
    cur = p_dict["1.2"]
    objs = [
        "• Aplicar un instrumento de diagnóstico de cultura basado en el Competing Values Framework (Cameron & Quinn / OCAI) adaptado a la Municipalidad Provincial de Tacna.",
        "• Determinar el perfil de cultura actual y deseada a 5 años (2030) y calcular la brecha matemática neta por dimensión.",
        "• Identificar los supuestos básicos subyacentes mediante el análisis hermenéutico de artefactos (ROF 2022, RISST 2024, PEI 2025-2030, Ley 27815) y de discurso institucional.",
        "• Detectar la brecha entre valores adoptados formalmente por la entidad y la conducta observada en la gestión pública cotidiana.",
        "• Derivar las implicancias del perfil cultural para condicionar la hoja de ruta y la estrategia de gestión del cambio del PETI.",
        "• Formular los valores organizacionales en términos estrictamente conductuales y redactar formalmente la Sección 2.3 y Sección 2.4 del plan."
    ]
    for ob in objs:
        cur = insert_p_after(cur, ob, style="Normal", space_after=3)

    # 1.3 Tiempo de duración
    cur = p_dict["1.3"]
    cur = insert_p_after(cur, "100 minutos (Sesión 2 en laboratorio · evaluación procedimental).", style="Normal", space_after=6)

    # 1.4 Resultados de aprendizaje
    cur = p_dict["1.4"]
    ras = [
        "• Evalúa y cuantifica la cultura organizacional mediante marcos analíticos validados (CVF/OCAI) aplicados al contexto del sector público subnacional.",
        "• Deduce supuestos básicos invisibles que condicionan el comportamiento institucional a partir de reglamentos, manuales de organización y funciones.",
        "• Formula declaraciones de valor exigibles operacionalizadas como conductas medibles con consecuencias disciplinarias y renuncias deliberadas.",
        "• Diseña estrategias de viabilización y mitigación cultural para proyectos de Gobierno Digital y Transformación Digital dentro del PETI."
    ]
    for ra in ras:
        cur = insert_p_after(cur, ra, style="Normal", space_after=3)

    # 1.5 Recursos (Table 0)
    t0 = doc.tables[0]
    recursos_data = [
        ("Python & Bibliotecas Científicas", "3.14.5 / pandas 3.0, numpy 2.4", "Control de calidad de datos, procesamiento estadístico y cálculo de brechas culturales."),
        ("Matplotlib", "3.11.2", "Generación y renderizado en alta resolución (300 DPI) del gráfico de radar comparativo."),
        ("OpenPyXL", "3.1.5", "Generación de libros de cálculo Excel con resultados y supuestos (Anexos B y D)."),
        ("Repositorio GitHub Oficial", "Git 2.x / Tag taller-05", f"Control de versiones y evidencias enlazadas: {GH_TAG}"),
        ("PEI 2025-2030 (MPT)", "Aprobado 2025 (135 pp.)", "Fuente institucional para analizar la Misión, OEI.07 (Gobierno Digital) y AEI.07.01-04."),
        ("Organigrama ROF 2022 (MPT)", "Edición oficial 2022", "Mapeo de la estructura orgánica edilicia, gerencias de línea y posición de OGTIC."),
        ("RISST MPT", "Versión 02 (nov. 2024, 51 pp.)", "Artefacto institucional para inferir supuestos de rigidez jerárquica y enfoque punitivo."),
        ("Ley N° 27815 (Código de Ética)", "Vigente", "Marco deontológico de la función pública para contrastar valores adoptados y conductas.")
    ]

    while len(t0.rows) > 1:
        tr = t0.rows[-1]._tr
        tr.getparent().remove(tr)

    format_row(t0.rows[0], "16285C", is_header=True, font_size=8.5, font_color=RGBColor(255, 255, 255))
    for r_item in recursos_data:
        new_row = t0.add_row()
        new_row.cells[0].paragraphs[0].text = r_item[0]
        new_row.cells[1].paragraphs[0].text = r_item[1]
        new_row.cells[2].paragraphs[0].text = r_item[2]
        format_row(new_row, "F8FAFC" if len(t0.rows)%2==1 else "FFFFFF", font_size=8)
        new_row.cells[0].paragraphs[0].runs[0].font.bold = True

    set_table_borders(t0, color="CBD5E1")

    # 1.6 Seguridad
    cur = p_dict["1.6"]
    for p in doc.paragraphs:
        if "Reglas de uso del laboratorio y de alcance que se respetaron" in p.text:
            p.text = (
                "Durante la ejecución del diagnóstico y el trabajo en laboratorio se observaron estrictamente las siguientes medidas de seguridad física, lógica y ética:\n"
                "1. Anonimato riguroso conforme a la Ley N° 29733 (Protección de Datos Personales) y D.S. 016-2024-JUS: El instrumento no recolectó ningún identificador personal (nombres, DNI ni correos personales). Las unidades orgánicas se agregaron únicamente para gerencias con al menos cinco respondientes, blindando la identidad de los trabajadores.\n"
                "2. Agregación y confidencialidad: Los resultados se entregan consolidados por dimensiones institucionales; jamás se divulgan microdatos individuales.\n"
                "3. Neutralidad técnica deontológica: El equipo formulador mantuvo una perspectiva descriptiva-propositiva, absteniéndose de emitir juicios de valor sobre autoridades o funcionarios.\n"
                "4. Custodia y destrucción de datos: Las respuestas digitales permanecen almacenadas en entornos cifrados locales y serán eliminadas al concluir el semestre académico.\n"
                "5. Normas de uso del centro de cómputo EPIS-UPT: Uso adecuado de los recursos computacionales, ejecución de scripts sin privilegios administrativos invasivos y respeto a la política de red institucional."
            )
            p.style = "Normal"
            p.paragraph_format.space_after = Pt(8)

    # 4. RELLENAR SECCIÓN 2: PROCEDIMIENTO O METODOLOGÍA (SIN PASO G Y SIN CUADROS REDUNDANTES)
    p_paso_a = p_dict["Paso A"]
    p_paso_b = p_dict["Paso B"]
    p_paso_c = p_dict["Paso C"]

    p_paso_b._p.getparent().remove(p_paso_b._p)
    p_paso_c._p.getparent().remove(p_paso_c._p)

    p_paso_a.text = "Paso A — Aplicar el Instrumento del Competing Values Framework (15 min)"
    p_paso_a.style = "Heading 2"

    cur = p_paso_a
    cur = insert_p_after(cur, "Determinar el perfil cultural de la Municipalidad Provincial de Tacna mediante la aplicación del instrumento OCAI de seis dimensiones con escala de suma forzada (100 puntos) entre los cuadrantes Clan (A), Adhocracia (B), Mercado (C) y Jerarquía (D) en dos momentos temporales: Actual y Deseado a 5 años (2030).", bold_prefix="• Qué se buscaba: ")
    cur = insert_p_after(cur, f"Diseño del instrumento contextualizado '02_identidad/CU01_instrumento.md' y levantamiento de una muestra de N=35 encuestas distribuidas en seis gerencias representativas de la sede central de la MPT, registradas en '02_identidad/encuesta_cultura.csv' (versionado en GitHub: {GH_TAG}/02_identidad/encuesta_cultura.csv).", bold_prefix="• Comando o acción: ")
    cur = insert_p_after(cur, f"Archivo de especificación del instrumento '02_identidad/CU01_instrumento.md' con las 6 dimensiones adaptadas a la corporación edil y dataset estructurado '02_identidad/encuesta_cultura.csv' conteniendo 35 registros y 51 variables de evaluación.", bold_prefix="• Evidencia de que funcionó: ")

    # Paso B
    cur = insert_p_after(cur, "Paso B — Calcular el Perfil Cultural y Brechas (15 min)", style="Heading 2")
    cur = insert_p_after(cur, "Ejecutar el control de calidad matemático estricto (exclusión de encuestas cuya suma en cualquier bloque no sea exactamente 100), calcular promedios dimensionales y globales, determinar brechas netas, identificar culturas dominantes, evaluar congruencia y generar la visualización en radar.", bold_prefix="• Qué se buscaba: ")
    cur = insert_p_after(cur, f"Desarrollo y ejecución del script automatizado en Python: 'python 02_identidad/CU02_perfil_cultura.py' (código: {GH_TAG}/02_identidad/CU02_perfil_cultura.py).", bold_prefix="• Comando o acción: ")
    cur = insert_p_after(cur, (
        "El script auditó las 35 encuestas, identificando y excluyendo con éxito 3 casos inconsistentes (R015 en OGAF con 110 pts; R024 en OGACGD con 90 pts; R031 en GTPSC con 105 pts), procesando 32 encuestas válidas (91.4% de efectividad). Se generó la matriz consolidada '02_identidad/CU02_perfil_cultura.csv' y el reporte de calidad 'docs/evidencias/S05/salidas/control_calidad.txt'.\n\n"
        "Síntesis de puntajes globales calculados:\n"
        "• Jerarquía (Control): Actual 44.10 pts | Deseado 19.04 pts | Brecha: -25.06 pts (Cultura Dominante Actual)\n"
        "• Adhocracia (Innovación): Actual 12.62 pts | Deseado 34.14 pts | Brecha: +21.52 pts (Cultura Dominante Deseada al 2030)\n"
        "• Clan (Colaboración): Actual 17.56 pts | Deseado 28.70 pts | Brecha: +11.14 pts\n"
        "• Mercado (Resultados): Actual 25.72 pts | Deseado 18.12 pts | Brecha: -7.60 pts"
    ), bold_prefix="• Evidencia de que funcionó: ")

    # Inserción de imagen
    if os.path.exists("02_identidad/CU_perfil_cultura.png"):
        p_img = insert_p_after(cur, "")
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(2)
        r_img = p_img.add_run()
        r_img.add_picture("02_identidad/CU_perfil_cultura.png", width=Inches(4.6))
        p_cap = insert_p_after(p_img, "Figura 2.1: Radar del Perfil de Cultura Organizacional en la Municipalidad Provincial de Tacna (Actual vs Deseado 2030)")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True
        p_cap.runs[0].font.color.rgb = RGBColor(71, 85, 105)
        cur = p_cap

    # Paso C
    cur = insert_p_after(cur, "Paso C — Identificar los Supuestos Básicos (10 min)", style="Heading 2")
    cur = insert_p_after(cur, "Descubrir los supuestos inconscientes que rigen las decisiones operativas reales mediante la confrontación sistemática entre artefactos físicos/documentales y el discurso de valores adoptados.", bold_prefix="• Qué se buscaba: ")
    cur = insert_p_after(cur, f"Análisis hermenéutico y de discurso sobre el ROF 2022, RISST 2024 (51 pp.), PEI 2025-2030 y Ley 27815, sintetizado en '02_identidad/CU03_supuestos.csv' (URL: {GH_TAG}/02_identidad/CU03_supuestos.csv) y en el Anexo D (Excel).", bold_prefix="• Comando o acción: ")
    cur = insert_p_after(cur, (
        "Se identificaron 6 supuestos profundos, de los cuales 5 evidencian una contradicción flagrante con los valores formales:\n"
        "1. Exigencia de 4 a 7 firmas correlativas físicas vs 'Modernización': Revela el supuesto de que la firma en papel con sello de tinta es el único blindaje legal ante el OCI.\n"
        "2. Bases de datos tributarias y catastrales aisladas vs 'Vocación de servicio': Revela el supuesto de que la información es patrimonio y fuente de poder del área.\n"
        "3. Servidores legados sin contingencia vs 'Innovación tecnológica': Revela el supuesto de que si el sistema funciona hoy, no se actualiza para evitar auditorías.\n"
        "4. RISST punitivo y subreporte de incidentes vs 'Mejora continua': Revela el supuesto de que reportar un error acarrea castigo disciplinario.\n"
        "5. OGTIC relegada a cotizar hardware vs 'Liderazgo digital': Revela el supuesto de que TI es solo soporte operativo de mantenimiento de PCs.\n"
        "6. Publicidad de sesiones de Concejo vs 'Transparencia': Coincide plenamente con el imperio de la ley."
    ), bold_prefix="• Evidencia de que funcionó: ")

    # Paso D
    cur = insert_p_after(cur, "Paso D — Derivar Implicancias para el PETI (10 min)", style="Heading 2")
    cur = insert_p_after(cur, "Traducir los hallazgos culturales cuantitativos y los supuestos básicos inferidos en riesgos de implantación y definir las estrategias técnicas y de cambio organizacional que deben incorporarse en el PETI.", bold_prefix="• Qué se buscaba: ")
    cur = insert_p_after(cur, f"Construcción de la matriz estratégica '02_identidad/CU04_implicancias_peti.md' ({GH_TAG}/02_identidad/CU04_implicancias_peti.md), articulándola con los objetivos de Gobierno Digital del PEI 2025-2030 (OEI.07).", bold_prefix="• Comando o acción: ")
    cur = insert_p_after(cur, (
        "Definición de salvaguardas que condicionan la hoja de ruta (Sección 7.3) y el plan de gestión del cambio (Sección 10):\n"
        "• Ante la Jerarquía y temor a OCI: Blindaje normativo del Sistema de Gestión Documental Cero Papel mediante Directiva con valor probatorio legal pleno de la Firma Digital RENIEC.\n"
        "• Ante el feudalismo de datos: Programa de Gobernanza de Datos con designación formal de Custodios previo a la integración del Catastro Multipropósito.\n"
        "• Ante la cultura punitiva: Política de 'Reporte Libre de Culpa' (Blameless Post-Mortem) para incidentes de ciberseguridad.\n"
        "• Ante la divergencia interdepartamental: Estrategia de comunicación segmentada por perfil cultural de cada gerencia."
    ), bold_prefix="• Evidencia de que funcionó: ")

    # Paso E
    cur = insert_p_after(cur, "Paso E — Formular los Valores y Redactar Secciones 2.3 y 2.4 (10 min)", style="Heading 2")
    cur = insert_p_after(cur, "Formular el marco de valores institucionales en enunciados conductuales con consecuencias disciplinarias y renuncias expresas, y redactar de forma integral las Secciones 2.3 y 2.4 del PETI de la MPT.", bold_prefix="• Qué se buscaba: ")
    cur = insert_p_after(cur, f"Redacción de '02_identidad/2.3_valores.md' ({GH_TAG}/02_identidad/2.3_valores.md) y '02_identidad/2.4_cultura.md' ({GH_TAG}/02_identidad/2.4_cultura.md), incorporando criterios de decisión tecnológica.", bold_prefix="• Comando o acción: ")
    cur = insert_p_after(cur, (
        "Cinco valores institucionales conductuales operativos formulados (Integridad y probidad, Vocación de servicio y empatía, Innovación abierta y mejora continua, Transparencia de datos, Colaboración interdisciplinaria) con sus conductas que cumplen, violan, sanciones y renuncias explícitas, así como cuatro criterios vinculantes para resolver tensiones de TI (sistemas interoperables abiertos vs propietarios cerrados, seguridad vs disponibilidad)."
    ), bold_prefix="• Evidencia de que funcionó: ")

    # Paso F
    cur = insert_p_after(cur, "Paso F — Validar y Corregir (25 min)", style="Heading 2")
    cur = insert_p_after(cur, "Someter los resultados a tres comprobaciones rigurosas y corregir cualquier desviación metodológica antes del cierre de sesión.", bold_prefix="• Qué se buscaba: ")
    cur = insert_p_after(cur, "Auditoría estadística del dataset, balance dimensional y contrastación empírica con decisiones de gobierno de la MPT.", bold_prefix="• Comando o acción: ")
    cur = insert_p_after(cur, (
        "Las tres comprobaciones resultaron plenamente satisfactorias:\n"
        "1. Comprobación de Muestra multi-área: Cumplida. Participaron servidores de OGTIC (n=6), GGT (n=7), OGAF (n=6), OGACGD (n=5), GTPSC (n=5) y OGPPMI (n=3).\n"
        "2. Comprobación de Suma 100 por bloque: Cumplida. El algoritmo de control de calidad descartó las 3 encuestas erróneas, garantizando que el 100% de la muestra analizada sume exactamente 100 puntos.\n"
        "3. Comprobación de Contrastación con la realidad: Cumplida. La brecha extrema en Adhocracia (+21.52) coincide con la reciente aprobación del PEI 2025-2030, donde la Alta Dirección incorporó formalmente el OEI.07 ('Fortalecer el Gobierno Digital') con metas de automatización masiva de trámites (AEI.07.02)."
    ), bold_prefix="• Evidencia de que funcionó: ")

    # 5. RELLENAR SECCIÓN 3: RESULTADOS (TABLE 1 OFICIAL RELLENADA AL 100%)
    # En la plantilla, Table 1 es la tabla de resultados oficial
    t1 = doc.tables[1]
    
    chk_14 = [
        ("1", "Instrumento del CVF aplicado con las seis dimensiones y los dos momentos", "Logrado", f"{GH_TAG}/02_identidad/CU01_instrumento.md"),
        ("2", "Al menos 10 respondientes o el 30% del personal de múltiples áreas", "Logrado", f"N=35 en {GH_TAG}/02_identidad/encuesta_cultura.csv"),
        ("3", "Control de calidad ejecutado (exclusión y reporte de sumas ≠ 100)", "Logrado", f"{GH_TAG}/docs/evidencias/S05/salidas/control_calidad.txt"),
        ("4", "Perfil de cultura actual y deseada calculado con la brecha por tipo", "Logrado", f"{GH_TAG}/02_identidad/CU02_perfil_cultura.csv"),
        ("5", "Cultura dominante y cultura deseada identificadas con su puntaje", "Logrado", "Actual: Jerarquía (44.10) | Deseada: Adhoc (34.14)"),
        ("6", "Análisis de congruencia por dimensión con veredicto", "Logrado", "Veredicto: ALTA CONGRUENCIA (6 dimensiones)"),
        ("7", "Perfil por área funcional (para áreas con N ≥ 5 respondientes)", "Logrado", f"{GH_TAG}/docs/evidencias/S05/salidas/analisis_perfil.txt"),
        ("8", "Gráfico de radar generado en alta resolución (300 DPI)", "Logrado", f"{GH_TAG}/02_identidad/CU_perfil_cultura.png"),
        ("9", "Al menos 5 supuestos básicos inferidos de artefactos con evidencia", "Logrado", f"6 supuestos en {GH_TAG}/02_identidad/CU03_supuestos.csv"),
        ("10", "Al menos 3 brechas entre valor adoptado y supuesto básico", "Logrado", f"5 brechas en {GH_TAG}/02_identidad/CU03_supuestos.csv"),
        ("11", "Tabla de implicancias con estrategia de implantación por hallazgo", "Logrado", f"{GH_TAG}/02_identidad/CU04_implicancias_peti.md"),
        ("12", "Valores formulados como conductas con consecuencias y renuncias", "Logrado", f"5 valores en {GH_TAG}/02_identidad/2.3_valores.md"),
        ("13", "Al menos un valor que gobierne una decisión tecnológica concreta", "Logrado", "4 tensiones y criterios TI en Sección 2.3.2"),
        ("14", "Secciones 2.3 y 2.4 redactadas formalmente con etiquetas Git", "Logrado", f"Tags v0.5 y taller-05 en {GH_BASE}/tags")
    ]

    # Vaciar filas anteriores excepto encabezado
    while len(t1.rows) > 1:
        tr = t1.rows[-1]._tr
        tr.getparent().remove(tr)

    format_row(t1.rows[0], "16285C", is_header=True, font_size=8.5, font_color=RGBColor(255, 255, 255))
    t1.cell(0, 0).width = Inches(0.4)
    t1.cell(0, 1).width = Inches(2.7)
    t1.cell(0, 2).width = Inches(0.9)
    t1.cell(0, 3).width = Inches(2.9)

    for item in chk_14:
        new_row = t1.add_row()
        new_row.cells[0].paragraphs[0].text = item[0]
        new_row.cells[1].paragraphs[0].text = item[1]
        new_row.cells[2].paragraphs[0].text = item[2]
        new_row.cells[3].paragraphs[0].text = item[3]
        format_row(new_row, "F8FAFC" if len(t1.rows)%2==1 else "FFFFFF", font_size=7.5)
        new_row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        new_row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        new_row.cells[2].paragraphs[0].runs[0].font.bold = True
        new_row.cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(21, 128, 61)

    set_table_borders(t1, color="CBD5E1")

    for p in doc.paragraphs:
        if "Si algo no se logró, explícalo aquí" in p.text:
            p.text = (
                f"Evaluación global de resultados: Todos los catorce resultados previstos en la lista de comprobación de la Semana 05 fueron alcanzados satisfactoriamente sin observaciones pendientes. Todo el código fuente, datasets, gráficos y matrices se encuentran versionados y verificables en el repositorio oficial de GitHub: {GH_TAG}."
            )
            p.style = "Normal"
            p.paragraph_format.space_after = Pt(8)

    # 6. RELLENAR SECCIÓN 4: CONCLUSIONES
    for p in doc.paragraphs:
        if "Mínimo tres. Una conclusión no resume lo que hiciste" in p.text:
            p.text = (
                "1. Supremacía de los Supuestos Básicos sobre el Discurso Formal:\n"
                "El diagnóstico demostró de forma concluyente que cuando los valores adoptados declarados (modernización, celeridad y servicio al vecino) colisionan con los supuestos básicos subyacentes (el miedo al OCI, la necesidad de sellos físicos y la propiedad patrimonial de los datos), gobiernan indefectiblemente los supuestos. Un Plan Estratégico de TI que ignore esta tensión comete el error fatal de diseñar sistemas para una corporación idílica inexistente, garantizando el rechazo pasivo de los usuarios.\n\n"
                "2. El Perfil Cultural como Condicionante de la Estrategia de Despliegue:\n"
                "El perfil cultural medido determina la metodología de presentación y la secuencia de implantación que debe adoptarse para cada proyecto tecnológico. En una entidad con cultura dominante de Jerarquía (44.10 puntos) como la MPT, un proyecto como el Sistema de Gestión Documental Cero Papel no puede sostenerse únicamente en la 'innovación'; requiere previamente resoluciones de Gerencia Municipal que otorguen pleno valor probatorio y blindaje legal a la firma digital. Asimismo, la divergencia entre gerencias exige presentar el plan con métricas de recaudación a Gestión Tributaria y con garantías de control a Administración.\n\n"
                "3. Exigibilidad de los Valores mediante la Renuncia Explícita en Decisiones de TI:\n"
                "Un valor que no implica renunciar a nada no orienta ninguna decisión real ni resuelve conflictos operativos. Formular los valores organizacionales de la MPT como conductas observables (cumplimiento vs violación), sanciones disciplinarias y renuncias deliberadas (renunciar a la discrecionalidad, al 'favor político' y al feudalismo de la información) es lo que los vuelve exigibles. Esto permite derivar políticas vinculantes para el PETI, tales como la obligatoriedad de arquitecturas interoperables de estándares abiertos y la adopción de una cultura de reporte libre de culpa (Blameless Post-Mortem) en ciberseguridad."
            )
            p.style = "Normal"
            p.paragraph_format.space_after = Pt(8)

    # 7. RELLENAR SECCIÓN 5: CUESTIONARIO
    for p in doc.paragraphs:
        if "Copia cada pregunta de la guía de la semana y respóndela debajo" in p.text:
            p.text = (
                "Pregunta 1 (Pregunta de transferencia de la guía · Paso G):\n"
                "¿Qué riesgo correría una organización real si este diagnóstico de cultura se hiciera mal o se omitiera?\n"
                "Respuesta: Correría el riesgo de invertir tiempo y presupuesto en soluciones tecnológicas que la organización terminará saboteando o subutilizando. En una institución pública como la Municipalidad Provincial de Tacna, omitir el diagnóstico significaría asumir erróneamente que el personal está listo para operar en un entorno digital ágil. En la práctica, al persistir el supuesto arraigado de que 'la firma en papel es el único blindaje ante la Contraloría', los servidores imprimirían digitalizaciones para sellarlas manualmente, duplicando los tiempos de trámite y convirtiendo el sistema informático en un elefante blanco burocrático.\n\n"
                "Pregunta 2:\n"
                "¿Por qué el Competing Values Framework emplea una escala de suma forzada (100 puntos) en lugar de una escala Likert convencional?\n"
                "Respuesta: Porque las escalas Likert tradicionales sufren de sesgo de deseabilidad social, permitiendo que el encuestado califique todas las dimensiones como 'muy importantes', lo que oculta las tensiones reales de la organización. La suma forzada de 100 puntos obliga matemáticamente a elegir y priorizar entre valores en competencia (flexibilidad vs control, enfoque interno vs externo), reflejando con exactitud los trade-offs de gestión que ocurren en la vida real de la entidad.\n\n"
                "Pregunta 3:\n"
                "¿Cómo condiciona la cultura dominante de Jerarquía (44.10 pts) la implantación del Sistema de Gestión Documental y Expediente Digital (SGD) en la MPT?\n"
                "Respuesta: Condiciona a que la implantación no puede ser un simple despliegue técnico de software, sino que debe estar precedida por un hito normativo de blindaje legal. Se requiere emitir Directivas Municipales aprobadas por Resolución de Gerencia Municipal y Decretos de Alcaldía que establezcan formalmente la equivalencia probatoria de la firma digital (RENIEC) frente a la firma manuscrita, eliminando el temor a responsabilidades ante el Órgano de Control Institucional (OCI).\n\n"
                "Pregunta 4:\n"
                "¿Cuál es el rol de explicitar a qué se renuncia al formular un valor organizacional y cómo influye en las decisiones tecnológicas?\n"
                "Respuesta: La renuncia explícita define el costo y el compromiso real de sostener el valor. Si un valor no declara a qué se renuncia, queda reducido a un lema propagandístico sin efecto normativo. Al explicitar que para sostener la 'Transparencia de Datos' se renuncia a que las gerencias retengan las bases de datos como cotos de poder, se habilita una política vinculante de TI: se prohíbe la compra de sistemas propietarios cerrados y se impone la interoperabilidad transversal obligatoria."
            )
            p.style = "Normal"
            p.paragraph_format.space_after = Pt(8)

    # 8. RELLENAR SECCIÓN 6: REFERENCIAS BIBLIOGRÁFICAS
    for p in doc.paragraphs:
        if "Normas técnicas, marcos profesionales, documentación oficial y bibliografía indexada" in p.text:
            p.text = (
                "Cameron, K. S. y Quinn, R. E. (2011). Diagnosing and Changing Organizational Culture: Based on the Competing Values Framework (3.ª ed.). Jossey-Bass.\n\n"
                "Congreso de la República del Perú. (2002). Ley N° 27815: Ley del Código de Ética de la Función Pública. Diario Oficial El Peruano.\n\n"
                "Congreso de la República del Perú. (2003). Ley N° 27972: Ley Orgánica de Municipalidades. Diario Oficial El Peruano.\n\n"
                "Congreso de la República del Perú. (2011). Ley N° 29733: Ley de Protección de Datos Personales y su reglamento modificado por D. S. 016-2024-JUS.\n\n"
                "González Millán, J. (2020). Manual práctico de planeación estratégica. Ediciones Díaz de Santos.\n\n"
                "ISACA. (2018). COBIT 2019 Framework: Governance and Management Objectives (Componente «Cultura, Ética y Comportamiento»). Information Systems Audit and Control Association.\n\n"
                "Kotter, J. P. (2012). Leading Change. Harvard Business Review Press.\n\n"
                "López Posada, L. M. (2016). Cultura organizacional: entre el individualismo y el colectivismo. Sello Editorial Universidad del Tolima.\n\n"
                "Municipalidad Provincial de Tacna. (2022). Reglamento de Organización y Funciones (ROF 2022) y Organigrama Estructural Institucional. MPT.\n\n"
                "Municipalidad Provincial de Tacna. (2024). Reglamento Interno de Seguridad y Salud en el Trabajo (RISST-001, Versión 02). Aprobado el 14/11/2024.\n\n"
                "Municipalidad Provincial de Tacna. (2025). Plan Estratégico Institucional de la Municipalidad Provincial de Tacna 2025-2030 (PEI 2025-2030). MPT.\n\n"
                "Rodríguez Bermúdez, J. R. (2015). Usos estratégicos de las TIC. Editorial UOC.\n\n"
                "Schein, E. H. y Schein, P. (2017). Organizational Culture and Leadership (5.ª ed.). Wiley."
            )
            p.style = "Normal"
            p.paragraph_format.space_after = Pt(8)

    # 9. RELLENAR SECCIÓN 7: ANEXOS
    for p in doc.paragraphs:
        if "Capturas completas, archivos de configuración y salidas extensas" in p.text:
            p.text = (
                f"Anexo A: Instrumento de Diagnóstico CVF (OCAI) · Municipalidad Provincial de Tacna\n"
                f"Archivo: anexos/anexo_A_instrumento_cvf.pdf (y .docx) | URL: {GH_TAG}/anexos/anexo_A_instrumento_cvf.pdf\n"
                f"Descripción: Cuestionario completo de seis dimensiones, con las cuatro afirmaciones contextualizadas a la corporación edil, regla de suma 100 e instrucciones metodológicas.\n\n"
                f"Anexo B: Matriz Consolidada de Resultados de Cultura Organizacional\n"
                f"Archivo: anexos/anexo_B_resultados_cultura.xlsx | URL: {GH_TAG}/anexos/anexo_B_resultados_cultura.xlsx\n"
                f"Descripción: Libro en formato Microsoft Excel con las 32 encuestas conformes depuradas y las hojas de promedios, brechas dimensionales y datos del radar.\n\n"
                f"Anexo C: Gráfico de Radar del Perfil Cultural (Alta Resolución)\n"
                f"Archivo: anexos/anexo_C_perfil_cultura.png | URL: {GH_TAG}/anexos/anexo_C_perfil_cultura.png\n"
                f"Descripción: Imagen renderizada a 300 DPI que ilustra el contraste polar entre la cultura Actual (Jerarquía dominante) y la cultura Deseada al 2030 (Adhocracia).\n\n"
                f"Anexo D: Matriz de Supuestos Básicos e Inferencias de la MPT\n"
                f"Archivo: anexos/anexo_D_supuestos_basicos.xlsx | URL: {GH_TAG}/anexos/anexo_D_supuestos_basicos.xlsx\n"
                f"Descripción: Matriz en Excel estructurada con los seis supuestos básicos inferidos a partir del ROF 2022, RISST 2024 y PEI 2025-2030, detallando las brechas observadas.\n\n"
                f"Anexo E: Secciones Oficiales 2.3 y 2.4 del PETI de la MPT\n"
                f"Archivo: anexos/anexo_E_secciones_2_3_2_4.pdf (y .docx) | URL: {GH_TAG}/anexos/anexo_E_secciones_2_3_2_4.pdf\n"
                f"Descripción: Texto normativo final aprobado de la Sección 2.3 (Valores institucionales y tecnológicos) y Sección 2.4 (Diagnóstico cultural), versionado bajo la etiqueta v0.5."
            )
            p.style = "Normal"
            p.paragraph_format.space_after = Pt(8)

    doc.save(output_docx)
    print(f"Documento Word oficial guardado con éxito: {output_docx}")

    # 10. ABRIR CON WORD COM, ACTUALIZAR TOC Y EXPORTAR A PDF
    print("\nActualizando campos, tabla de contenidos y exportando a PDF vía Microsoft Word...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False

    abs_docx = os.path.abspath(output_docx)
    abs_pdf = os.path.abspath(output_pdf)

    try:
        w_doc = word.Documents.Open(abs_docx)
        w_doc.Fields.Update()
        try:
            for toc in w_doc.TablesOfContents:
                toc.Update()
        except Exception as e_toc:
            print("Aviso al actualizar TOC:", e_toc)

        w_doc.Save()
        w_doc.SaveAs(abs_pdf, FileFormat=17)
        w_doc.Close()
        print(f"Exportación a PDF completada exitosamente: {output_pdf}")
    except Exception as e:
        print("Error en Word COM:", e)
    finally:
        word.Quit()

if __name__ == "__main__":
    run()
